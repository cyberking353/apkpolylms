import os
from functools import wraps
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib import messages
from django.contrib.auth import logout
from django.utils import timezone
from django.db import connection, transaction
from django.db.models import Count, Avg, Q


def portal_required(view_func):
    @wraps(view_func)
    def wrapper(request, *args, **kwargs):
        if not request.user.is_authenticated or not request.user.is_staff:
            return redirect('accounts:login')
        return view_func(request, *args, **kwargs)
    return wrapper


def _is_htmx(request):
    """True when the request came from an htmx live-filter call."""
    return request.headers.get('HX-Request') == 'true'


# ── Auth ───────────────────────────────────────────────────────────────

def portal_login(request):
    """Admins now sign in through the same Portal-driven SSO as everyone else.
    This view just hands them off to the SSO handoff page (accounts:login),
    which itself redirects to the Portal."""
    if request.user.is_authenticated and request.user.is_staff:
        return redirect('portal:dashboard')
    return redirect('accounts:login')


def portal_logout(request):
    logout(request)
    return redirect('accounts:login')


# ── Dashboard ──────────────────────────────────────────────────────────

@portal_required
def dashboard(request):
    from accounts.models import Student, ACADEMIC_LEVEL_CHOICES
    from courses.models import Course, Enrollment, AcademicSettings
    from departments.models import Department
    from exams.models import Exam
    from datetime import date

    settings_obj = AcademicSettings.get_current()
    students = Student.objects.filter(role='student')
    lecturers = Student.objects.filter(role='lecturer')
    courses = Course.objects.all()

    stats = {
        'students':        students.count(),
        'lecturers':       lecturers.count(),
        'courses_live':    courses.filter(status='published').count(),
        'courses_total':   courses.count(),
        'courses_pending': courses.filter(status='pending').count(),
        'exams_pending':   Exam.objects.filter(exam_status='pending').count(),
        'enrollments':     Enrollment.objects.count(),
    }

    pending_courses = courses.filter(status='pending').select_related('lecturer', 'category')[:6]
    pending_exams   = Exam.objects.filter(exam_status='pending').select_related('course')[:6]
    recent_students = students.order_by('-date_joined')[:5]

    level_labels = dict(ACADEMIC_LEVEL_CHOICES)
    dept_map = dict(Department.objects.values_list('code', 'name'))

    dept_breakdown = [
        {'label': dept_map.get(d['department']) or d['department'] or 'Unassigned', 'count': d['count']}
        for d in students.values('department').annotate(count=Count('id')).order_by('-count')
    ]
    level_breakdown = [
        {'label': level_labels.get(l['academic_level']) or l['academic_level'] or 'Unassigned', 'count': l['count']}
        for l in students.values('academic_level').annotate(count=Count('id')).order_by('academic_level')
    ]

    import datetime as _dt
    today = _dt.date.today()

    return render(request, 'portal/dashboard.html', {
        'stats': stats,
        'pending_courses': pending_courses,
        'pending_exams': pending_exams,
        'recent_students': recent_students,
        'dept_breakdown': dept_breakdown,
        'level_breakdown': level_breakdown,
        'settings': settings_obj,
        'today': today,
    })


# ── Students ───────────────────────────────────────────────────────────

@portal_required
def student_list(request):
    from accounts.models import Student, get_department_choices, ACADEMIC_LEVEL_CHOICES
    qs = Student.objects.filter(role='student').order_by('last_name', 'first_name')
    dept  = request.GET.get('dept', '')
    level = request.GET.get('level', '')
    q     = request.GET.get('q', '')
    if dept:  qs = qs.filter(department=dept)
    if level: qs = qs.filter(academic_level=level)
    if q:     qs = qs.filter(Q(first_name__icontains=q)|Q(last_name__icontains=q)|Q(email__icontains=q)|Q(matric_number__icontains=q))
    ctx = {
        'students': qs, 'dept': dept, 'level': level, 'q': q,
        'dept_choices': get_department_choices(),
        'level_choices': ACADEMIC_LEVEL_CHOICES,
    }
    if _is_htmx(request):
        return render(request, 'portal/_student_table.html', ctx)
    return render(request, 'portal/student_list.html', ctx)


@portal_required
def student_export(request):
    from accounts.models import Student, get_department_choices, ACADEMIC_LEVEL_CHOICES

    if request.method == 'POST' or 'all' in request.GET:
        dept  = request.POST.get('dept', '') if request.method == 'POST' else ''
        level = request.POST.get('level', '') if request.method == 'POST' else ''

        qs = Student.objects.filter(role='student')
        if dept:  qs = qs.filter(department=dept)
        if level: qs = qs.filter(academic_level=level)

        if dept:
            qs = qs.order_by('last_name', 'first_name')
        else:
            qs = qs.order_by('department', 'last_name', 'first_name')

        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        from django.http import HttpResponse

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Students'
        header_font = Font(bold=True, color='FFFFFF', size=11)
        header_fill = PatternFill(start_color='1A6B72', end_color='1A6B72', fill_type='solid')
        headers = ['#', 'Surname', 'First Name', 'Email', 'Matric Number', 'Department', 'Level', 'Active', 'Date Joined']
        ws.append(headers)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center')

        for i, s in enumerate(qs, start=1):
            ws.append([
                i,
                s.last_name,
                s.first_name,
                s.email,
                s.matric_number or '',
                s.get_department_display() or '',
                s.get_academic_level_display() or '',
                'Yes' if s.is_active else 'No',
                s.date_joined.strftime('%d %b %Y') if s.date_joined else '',
            ])

        for col in ws.columns:
            max_len = max((len(str(c.value or '')) for c in col), default=0)
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        filename = f'students_{timezone.now():%Y%m%d_%H%M%S}.xlsx'
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        wb.save(response)
        return response

    dept_choices = [('', 'All Departments')] + list(get_department_choices())
    level_choices = [('', 'All Levels')] + list(ACADEMIC_LEVEL_CHOICES)
    return render(request, 'portal/student_export.html', {
        'dept_choices': dept_choices,
        'level_choices': level_choices,
    })


@portal_required
def student_create(request):
    from accounts.models import Student, get_department_choices, ACADEMIC_LEVEL_CHOICES
    if request.method == 'POST':
        p = request.POST
        pwd = p.get('password', '').strip()
        if not pwd:
            messages.error(request, 'Password is required.')
            return redirect('portal:student_create')
        if Student.objects.filter(email=p['email']).exists():
            messages.error(request, 'An account with this email already exists.')
        else:
            matric = (p.get('matric_number','') or '').strip() or None
            s = Student.objects.create_user(
                username=p['email'], email=p['email'],
                first_name=p.get('first_name',''), last_name=p.get('last_name',''),
                password=pwd,
                role='student',
                department=p.get('department',''),
                academic_level=p.get('academic_level',''),
                matric_number=matric,
                current_session=p.get('current_session',''),
                current_semester=p.get('current_semester',''),
                is_sms_synced=False,  # manual portal creation, not SMS-originated
            )
            from courses.models import Enrollment
            Enrollment.sync_for_student(s)
            messages.success(request, f'Student {s.get_full_name()} created.')
            return redirect('portal:student_list')
    return render(request, 'portal/student_form.html', {
        'is_new': True,
        'dept_choices': get_department_choices(),
        'level_choices': ACADEMIC_LEVEL_CHOICES,
    })


@portal_required
def student_edit(request, pk):
    from accounts.models import Student, get_department_choices, ACADEMIC_LEVEL_CHOICES
    student = get_object_or_404(Student, pk=pk, role='student')
    if request.method == 'POST':
        p = request.POST
        student.first_name    = p.get('first_name', student.first_name)
        student.last_name     = p.get('last_name', student.last_name)
        student.matric_number = (p.get('matric_number', '') or '').strip() or None
        student.department    = p.get('department', student.department)
        student.academic_level = p.get('academic_level', student.academic_level)
        student.is_active     = 'is_active' in p
        if p.get('password'):
            student.set_password(p['password'])
        student.save()
        from courses.models import Enrollment
        Enrollment.sync_for_student(student)
        messages.success(request, 'Student updated.')
        return redirect('portal:student_list')
    return render(request, 'portal/student_form.html', {
        'is_new': False, 'student': student,
        'dept_choices': get_department_choices(),
        'level_choices': ACADEMIC_LEVEL_CHOICES,
    })


@portal_required
def student_promote(request, pk):
    from accounts.models import Student
    student = get_object_or_404(Student, pk=pk, role='student')
    if student.promote():
        from courses.models import Enrollment
        Enrollment.sync_for_student(student)
        messages.success(request, f'{student.get_full_name()} promoted to {student.get_academic_level_display()}.')
    else:
        messages.warning(request, f'{student.get_full_name()} is already at the highest level (HND2).')
    return redirect('portal:student_list')


@portal_required
def student_reset_password(request, pk):
    """
    Quick, focused password reset for a student who's forgotten theirs — separate
    from the full Edit Student form since this is the one thing an admin needs to
    do fast, without wading through every other field. Since students don't self-
    register or receive email, the admin sets the new password directly and relays
    it to the student (in person, by phone, etc.).
    """
    from accounts.models import Student
    student = get_object_or_404(Student, pk=pk, role='student')
    if request.method == 'POST':
        new_password = request.POST.get('password', '').strip()
        if len(new_password) < 6:
            messages.error(request, 'Password must be at least 6 characters.')
            return render(request, 'portal/student_reset_password.html', {'student': student})
        student.set_password(new_password)
        student.save(update_fields=['password'])
        messages.success(
            request,
            f'Password reset for {student.get_full_name()}.'
        )
        return redirect('portal:student_list')
    return render(request, 'portal/student_reset_password.html', {'student': student})


# ── Lecturers ──────────────────────────────────────────────────────────

@portal_required
def lecturer_list(request):
    from accounts.models import Student, get_department_choices
    qs = Student.objects.filter(role='lecturer').order_by('last_name')
    q = request.GET.get('q', '').strip()
    dept = request.GET.get('dept', '')
    if q:
        qs = qs.filter(
            Q(first_name__icontains=q) | Q(last_name__icontains=q)
            | Q(email__icontains=q) | Q(staff_id__icontains=q))
    if dept:
        qs = qs.filter(department=dept)
    ctx = {
        'lecturers': qs, 'q': q, 'dept': dept,
        'dept_choices': get_department_choices(),
        'total_count': Student.objects.filter(role='lecturer').count(),
    }
    if _is_htmx(request):
        return render(request, 'portal/_lecturer_table.html', ctx)
    return render(request, 'portal/lecturer_list.html', ctx)


@portal_required
def lecturer_create(request):
    from accounts.models import Student, get_department_choices
    if request.method == 'POST':
        p = request.POST
        pwd = p.get('password', '').strip()
        if not pwd:
            messages.error(request, 'Password is required.')
            return redirect('portal:lecturer_create')
        if Student.objects.filter(email=p['email']).exists():
            messages.error(request, 'An account with this email already exists.')
        else:
            Student.objects.create_user(
                username=p['email'], email=p['email'],
                first_name=p.get('first_name',''), last_name=p.get('last_name',''),
                password=pwd,
                role='lecturer',
                department=p.get('department',''),
                staff_id=p.get('staff_id',''),
                title=p.get('title',''),
                is_active=True,
            )
            messages.success(request, 'Lecturer account created.')
            return redirect('portal:lecturer_list')
    return render(request, 'portal/lecturer_form.html', {
        'is_new': True,
        'dept_choices': get_department_choices(),
    })


@portal_required
def lecturer_edit(request, pk):
    from accounts.models import Student, get_department_choices
    lecturer = get_object_or_404(Student, pk=pk, role='lecturer')
    if request.method == 'POST':
        p = request.POST
        lecturer.title      = p.get('title', lecturer.title)
        lecturer.first_name = p.get('first_name', lecturer.first_name)
        lecturer.last_name  = p.get('last_name', lecturer.last_name)
        lecturer.staff_id   = p.get('staff_id', lecturer.staff_id)
        lecturer.department = p.get('department', lecturer.department)
        lecturer.is_active  = 'is_active' in p
        if p.get('password'):
            lecturer.set_password(p['password'])
        lecturer.save()
        messages.success(request, 'Lecturer updated.')
        return redirect('portal:lecturer_list')
    return render(request, 'portal/lecturer_form.html', {
        'is_new': False, 'lecturer': lecturer,
        'dept_choices': get_department_choices(),
    })


@portal_required
def lecturer_reset_password(request, pk):
    """Quick, focused password reset for a lecturer — see student_reset_password for the same idea."""
    from accounts.models import Student
    lecturer = get_object_or_404(Student, pk=pk, role='lecturer')
    if request.method == 'POST':
        new_password = request.POST.get('password', '').strip()
        if len(new_password) < 6:
            messages.error(request, 'Password must be at least 6 characters.')
            return render(request, 'portal/lecturer_reset_password.html', {'lecturer': lecturer})
        lecturer.set_password(new_password)
        lecturer.save(update_fields=['password'])
        messages.success(
            request,
            f'Password reset for {lecturer.display_name}.'
        )
        return redirect('portal:lecturer_list')
    return render(request, 'portal/lecturer_reset_password.html', {'lecturer': lecturer})


@portal_required
def lecturer_delete(request, pk):
    """Permanently delete a staff (lecturer) account. Their courses are kept
    (the course->lecturer link is SET_NULL) but the account itself is removed.
    Users are only ever managed here or via the SMS sync — this is a manual
    admin override."""
    from accounts.models import Student
    lecturer = get_object_or_404(Student, pk=pk, role='lecturer')
    if request.method == 'POST':
        if request.user.pk == lecturer.pk:
            messages.error(request, 'You cannot delete your own account.')
            return redirect('portal:lecturer_list')
        name = lecturer.display_name
        lecturer.delete()
        messages.success(request, f'{name} has been permanently deleted.')
    return redirect('portal:lecturer_list')


# ── Courses ────────────────────────────────────────────────────────────

@portal_required
def course_list(request):
    from courses.models import Course
    qs = Course.objects.select_related('lecturer', 'category').order_by('-created_at')
    status = request.GET.get('status', '')
    if status:
        qs = qs.filter(status=status)
    # Findability: free-text search across code/title/lecturer + level filter
    q = request.GET.get('q', '').strip()
    if q:
        from django.db.models import Q as _Q
        qs = qs.filter(
            _Q(code__icontains=q)
            | _Q(title__icontains=q)
            | _Q(lecturer__first_name__icontains=q)
            | _Q(lecturer__last_name__icontains=q)
            | _Q(lecturer__username__icontains=q)
        )
    level = request.GET.get('level', '')
    if level:
        qs = qs.filter(level=level)
    ctx = {
        'courses': qs, 'active_status': status,
        'status_choices': Course.STATUS_CHOICES,
        'q': q, 'active_level': level,
        'level_choices': Course.LEVEL_CHOICES,
        'total_count': Course.objects.count(),
    }
    if _is_htmx(request):
        return render(request, 'portal/_course_table.html', ctx)
    return render(request, 'portal/course_list.html', ctx)


def _youtube_embed_url(url):
    if not url:
        return ''
    import re
    patterns = [
        r'(?:youtube\.com/watch\?v=|youtu\.be/|youtube\.com/embed/)([a-zA-Z0-9_-]{11})',
    ]
    for p in patterns:
        m = re.search(p, url)
        if m:
            return f'https://www.youtube.com/embed/{m.group(1)}?modestbranding=1'
    return url


@portal_required
def course_review(request, pk):
    from courses.models import Course, CourseDepartment, Notification
    from accounts.models import get_department_choices
    from quizzes.models import Quiz
    course = get_object_or_404(Course, pk=pk)
    if request.method == 'POST':
        action = request.POST.get('action')
        note   = request.POST.get('review_note', '')
        if action == 'publish':
            depts = request.POST.getlist('departments')
            CourseDepartment.objects.filter(course=course).delete()
            for d in depts:
                CourseDepartment.objects.create(
                    course=course, department=d,
                    is_home=(d == course.home_department or not course.home_department)
                )
            course.publish(note=note)
            if course.lecturer:
                Notification.objects.create(
                    recipient=course.lecturer,
                    message=f'"{course.title}" has been published and is now live for students.',
                    link='/lecturer/',
                )
            messages.success(request, f'"{course.title}" published.')
        elif action == 'reject':
            course.reject(note=note)
            if course.lecturer:
                Notification.objects.create(
                    recipient=course.lecturer,
                    message=f'"{course.title}" was rejected. Review the admin note and resubmit.',
                    link=f'/lecturer/course/{course.slug}/edit/',
                )
            messages.warning(request, f'"{course.title}" rejected with note sent to lecturer.')
        return redirect('portal:course_list')
    modules = course.modules.prefetch_related('lectures').all()
    for m in modules:
        for lec in m.lectures.all():
            lec.embed_url = _youtube_embed_url(lec.video_url) if lec.video_url else ''
    quizzes = Quiz.objects.filter(course=course).prefetch_related(
        'questions__choices', 'module'
    ).order_by('module__week_number')
    exam = getattr(course, 'exam', None)
    if exam:
        from exams.models import ExamQuestion
        exam_questions = ExamQuestion.objects.filter(exam=exam).prefetch_related('choices').order_by('order')
    else:
        exam_questions = []
    return render(request, 'portal/course_review.html', {
        'course': course, 'modules': modules,
        'quizzes': quizzes, 'exam': exam, 'exam_questions': exam_questions,
        'dept_choices': get_department_choices(),
        'selected_depts': course.department_list,
    })


@portal_required
def course_delete(request, pk):
    from courses.models import (
        Course, CourseDepartment, Module, Lecture, LectureProgress,
        Enrollment, CourseMessage, CourseMessageReply,
    )
    from quizzes.models import Quiz, Question, Choice, QuizAttempt, Answer
    from exams.models import (
        Exam, ExamQuestion, ExamChoice, ExamAttempt, ExamAnswer,
        ExamViolation, ExamReschedule,
    )
    from results.models import Result

    course = get_object_or_404(Course, pk=pk)
    title = course.title

    # Only genuine academic RECORDS force archiving: grades and attempts are
    # history that cannot be rebuilt. Watch-progress and Q&A threads are not
    # enough to block deletion of an unwanted course.
    has_academic_data = (
        Result.objects.filter(course=course).exists()
        or QuizAttempt.objects.filter(quiz__course=course).exists()
        or ExamAttempt.objects.filter(exam__course=course).exists()
    )
    if has_academic_data:
        course.status = 'archived'
        course.save(update_fields=['status', 'updated_at'])
        messages.info(request,
            f'"{title}" has been archived. Its results, attempts and progress '
            f'are preserved; use Restore to bring it back.')
        return redirect('portal:course_list')

    # Delete deepest dependencies first. Django's CASCADE normally handles
    # this, but a production DB whose schema is behind the migrations can
    # still carry legacy foreign-key constraints without ON DELETE CASCADE,
    # which raises "FOREIGN KEY constraint failed" on a bare .delete().
    # For SQLite we suspend FK enforcement around the operation (the pragma
    # is a no-op inside a transaction, so it must run before it opens).
    fk_restored = False
    if connection.vendor == 'sqlite':
        cursor = connection.cursor()
        cursor.execute('PRAGMA foreign_keys')
        if cursor.fetchone()[0]:
            cursor.execute('PRAGMA foreign_keys = OFF')
            fk_restored = True
    try:
        with transaction.atomic():
            ExamAnswer.objects.filter(attempt__exam__course=course).delete()
            ExamViolation.objects.filter(attempt__exam__course=course).delete()
            ExamAttempt.objects.filter(exam__course=course).delete()
            ExamReschedule.objects.filter(exam__course=course).delete()
            ExamChoice.objects.filter(question__exam__course=course).delete()
            ExamQuestion.objects.filter(exam__course=course).delete()
            Exam.objects.filter(course=course).delete()

            Answer.objects.filter(question__quiz__course=course).delete()
            QuizAttempt.objects.filter(quiz__course=course).delete()
            Choice.objects.filter(question__quiz__course=course).delete()
            Question.objects.filter(quiz__course=course).delete()
            Quiz.objects.filter(course=course).delete()

            LectureProgress.objects.filter(lecture__module__course=course).delete()
            Lecture.objects.filter(module__course=course).delete()
            Module.objects.filter(course=course).delete()

            CourseMessageReply.objects.filter(thread__course=course).delete()
            CourseMessage.objects.filter(course=course).delete()
            Result.objects.filter(course=course).delete()
            Enrollment.objects.filter(course=course).delete()
            CourseDepartment.objects.filter(course=course).delete()

            course.delete()
    finally:
        if fk_restored:
            connection.cursor().execute('PRAGMA foreign_keys = ON')

    messages.success(request, f'"{title}" has been permanently deleted.')
    return redirect('portal:course_list')


@portal_required
def course_restore(request, pk):
    """Bring an archived course back as a draft so it can be edited/re-published."""
    from courses.models import Course
    course = get_object_or_404(Course, pk=pk)
    if course.status == 'archived':
        course.status = 'draft'
        course.save(update_fields=['status', 'updated_at'])
        messages.success(request, f'"{course.title}" has been restored to Draft.')
    return redirect('portal:course_list')


@portal_required
def course_create(request):
    from courses.models import Course, CourseDepartment, Category
    from accounts.models import Student, get_department_choices, ACADEMIC_LEVEL_CHOICES
    if request.method == 'POST':
        title = request.POST.get('title', '').strip()
        if not title:
            messages.error(request, 'Please give the course a title.')
            return render(request, 'portal/course_form.html', {
                'is_new': True, 'posted': request.POST,
                'departments': get_department_choices(),
                'selected_departments': request.POST.getlist('departments'),
                'levels': Course.LEVEL_CHOICES,
                'categories': Category.objects.all(),
                'lecturers': Student.objects.filter(role='lecturer', is_active=True).order_by('last_name'),
            })
        slug = _unique_slug_for_course(title)
        lecturer_id = request.POST.get('lecturer') or None
        course = Course.objects.create(
            title=title, slug=slug,
            category_id=request.POST.get('category') or None,
            description=request.POST.get('description', '').strip(),
            icon=request.POST.get('icon', '📚').strip() or '📚',
            level=request.POST.get('level', 'ND1'),
            semester=request.POST.get('semester', '1'),
            duration_hours=request.POST.get('duration_hours') or 0,
            lecturer_id=lecturer_id,
            status='published',
            instructor_name_override=request.POST.get('instructor_name_override', '').strip(),
            instructor_title_override=request.POST.get('instructor_title_override', '').strip(),
        )
        CourseDepartment.objects.filter(course=course).delete()
        for dept in request.POST.getlist('departments'):
            if dept:
                CourseDepartment.objects.create(
                    course=course, department=dept,
                    is_home=(dept == request.POST.get('home_department', ''))
                )
        messages.success(request, f'"{course.title}" created and published. You can review it and manage content from the course list.')
        return redirect('portal:course_list')

    return render(request, 'portal/course_form.html', {
        'is_new': True,
        'departments': get_department_choices(),
        'selected_departments': [],
        'levels': Course.LEVEL_CHOICES,
        'categories': Category.objects.all(),
        'lecturers': Student.objects.filter(role='lecturer', is_active=True).order_by('last_name'),
    })


@portal_required
def course_edit(request, pk):
    from courses.models import Course, CourseDepartment, Category, Module
    from accounts.models import Student, get_department_choices
    course = get_object_or_404(Course, pk=pk)
    if request.method == 'POST':
        course.title = request.POST.get('title', '').strip() or course.title
        course.category_id = request.POST.get('category') or None
        course.description = request.POST.get('description', '').strip()
        course.icon = request.POST.get('icon', course.icon).strip() or course.icon
        course.level = request.POST.get('level', course.level)
        course.semester = request.POST.get('semester', course.semester)
        course.duration_hours = request.POST.get('duration_hours') or 0
        course.lecturer_id = request.POST.get('lecturer') or None
        course.instructor_name_override = request.POST.get('instructor_name_override', '').strip()
        course.instructor_title_override = request.POST.get('instructor_title_override', '').strip()
        if 'status' in request.POST:
            course.status = request.POST['status']
        course.save()
        CourseDepartment.objects.filter(course=course).delete()
        for dept in request.POST.getlist('departments'):
            if dept:
                CourseDepartment.objects.create(
                    course=course, department=dept,
                    is_home=(dept == request.POST.get('home_department', ''))
                )
        messages.success(request, f'"{course.title}" updated.')
        return redirect('portal:course_edit', pk=pk)
    modules = Module.objects.filter(course=course).prefetch_related('lectures')
    return render(request, 'portal/course_form.html', {
        'is_new': False, 'course': course, 'modules': modules,
        'admin_mode': True,
        'departments': get_department_choices(),
        'selected_departments': course.department_list,
        'levels': Course.LEVEL_CHOICES,
        'categories': Category.objects.all(),
        'lecturers': Student.objects.filter(role='lecturer', is_active=True).order_by('last_name'),
    })


def _unique_slug_for_course(title):
    from django.utils.text import slugify
    from courses.models import Course
    base = slugify(title)
    slug = base
    n = 1
    while Course.objects.filter(slug=slug).exists():
        n += 1
        slug = f"{base}-{n}"
    return slug


# ── Exams ──────────────────────────────────────────────────────────────

@portal_required
def exam_list(request):
    from exams.models import Exam
    qs = Exam.objects.select_related('course').order_by('-course__created_at')
    status = request.GET.get('status', '')
    if status:
        qs = qs.filter(exam_status=status)
    q = request.GET.get('q', '').strip()
    if q:
        qs = qs.filter(
            Q(course__code__icontains=q) | Q(course__title__icontains=q)
            | Q(title__icontains=q))
    ctx = {
        'exams': qs, 'active_status': status, 'q': q,
        'total_count': Exam.objects.count(),
    }
    if _is_htmx(request):
        return render(request, 'portal/_exam_table.html', ctx)
    return render(request, 'portal/exam_list.html', ctx)


@portal_required
def exam_review(request, pk):
    from exams.models import Exam, ExamQuestion
    from courses.models import Notification
    exam = get_object_or_404(Exam, pk=pk)
    if request.method == 'POST':
        action = request.POST.get('action')
        note   = request.POST.get('review_note', '')
        if action == 'approve':
            from django.utils.dateparse import parse_datetime
            open_at  = parse_datetime(request.POST.get('open_at', ''))
            close_at = parse_datetime(request.POST.get('close_at', ''))
            if not open_at or not close_at:
                messages.error(request, 'Please set both open and close date/time.')
                return redirect('portal:exam_review', pk=pk)
            exam.time_limit_minutes = request.POST.get('time_limit_minutes') or exam.time_limit_minutes
            exam.pass_mark = request.POST.get('pass_mark') or exam.pass_mark
            exam.save()
            exam.approve(open_at=open_at, close_at=close_at, note=note)
            Notification.objects.create(
                recipient=exam.course.lecturer,
                message=f'Exam for "{exam.course.title}" has been approved and scheduled.',
                link=f'/lecturer/course/{exam.course.slug}/exam/manage/',
            )
            messages.success(request, f'Exam approved and scheduled.')
        elif action == 'reject':
            exam.exam_status = 'draft'
            exam.review_note = note
            exam.save(update_fields=['exam_status', 'review_note'])
            Notification.objects.create(
                recipient=exam.course.lecturer,
                message=f'Exam for "{exam.course.title}" was rejected. Check the admin note and resubmit.',
                link=f'/lecturer/course/{exam.course.slug}/exam/manage/',
            )
            messages.warning(request, 'Exam sent back to lecturer.')
        return redirect('portal:exam_list')
    questions = exam.questions.prefetch_related('choices').all()
    return render(request, 'portal/exam_review.html', {
        'exam': exam, 'questions': questions,
    })


# ── Results ────────────────────────────────────────────────────────────

@portal_required
def result_list(request):
    """
    Results explorer with two professional views instead of one endless table:
      view=course  (default) -> pick a course, see its gradebook + class stats
      view=student           -> pick a student, see their full report card
    """
    from results.models import Result
    from accounts.models import Student, get_department_choices, ACADEMIC_LEVEL_CHOICES
    from courses.models import Course, SEMESTER_CHOICES

    view  = request.GET.get('view', 'course')
    dept  = request.GET.get('dept', '')
    level = request.GET.get('level', '')
    yr    = request.GET.get('year', '')
    sem   = request.GET.get('sem', '')
    q     = request.GET.get('q', '').strip()
    year_choices = (Result.objects.values_list('academic_year', flat=True)
                    .distinct().order_by('-academic_year'))

    ctx = {
        'view': view, 'dept': dept, 'level': level,
        'yr': yr, 'sem': sem, 'q': q,
        'dept_choices': get_department_choices(),
        'level_choices': ACADEMIC_LEVEL_CHOICES,
        'year_choices': year_choices,
        'semester_choices': SEMESTER_CHOICES,
    }

    if view == 'student':
        sid = request.GET.get('student', '')
        if sid:
            student = get_object_or_404(Student, pk=sid, role='student')
            results = (Result.objects.filter(student=student)
                       .select_related('course').order_by('course__title'))
            if yr: results = results.filter(academic_year=yr)
            if sem: results = results.filter(semester=sem)
            totals = list(results.values_list('total', flat=True))
            ctx.update(
                student=student, results=results,
                avg_total=round(sum(totals) / len(totals), 1) if totals else None)
        else:
            students = (Student.objects.filter(role='student')
                        .annotate(n_results=Count('results'))
                        .filter(n_results__gt=0))
            if q:
                students = students.filter(
                    Q(first_name__icontains=q) | Q(last_name__icontains=q)
                    | Q(matric_number__icontains=q))
            if dept:  students = students.filter(department=dept)
            if level: students = students.filter(academic_level=level)
            ctx['students'] = (students
                               .order_by('last_name', 'first_name')[:200])
    else:
        cid = request.GET.get('course', '')
        if cid:
            course = get_object_or_404(Course, pk=cid)
            results = (Result.objects.filter(course=course)
                       .select_related('student').order_by('-total'))
            if dept:  results = results.filter(student__department=dept)
            if level: results = results.filter(student__academic_level=level)
            if yr:    results = results.filter(academic_year=yr)
            if sem:   results = results.filter(semester=sem)
            totals = list(results.values_list('total', flat=True))
            n = len(totals)
            ctx.update(
                course=course, results=results,
                stats={
                    'count': n,
                    'avg': round(sum(totals) / n, 1) if n else 0,
                    'high': max(totals) if n else 0,
                    'low': min(totals) if n else 0,
                    'pass_rate': round(100 * sum(1 for t in totals if t >= 50) / n) if n else 0,
                })
        else:
            courses = (Course.objects
                       .annotate(n_results=Count('results'),
                                 avg_total=Avg('results__total'))
                       .filter(n_results__gt=0))
            if q:
                courses = courses.filter(Q(code__icontains=q) | Q(title__icontains=q))
            if dept:  courses = courses.filter(coursedepartment__department=dept).distinct()
            if level: courses = courses.filter(level=level)
            if yr:    courses = courses.filter(results__academic_year=yr).distinct()
            if sem:   courses = courses.filter(semester=sem)
            ctx['courses'] = courses.order_by('code')[:200]

    if _is_htmx(request):
        return render(request, 'portal/_results_area.html', ctx)
    return render(request, 'portal/result_list.html', ctx)


@portal_required
def result_export(request):
    from results.models import Result
    from accounts.models import get_department_choices, ACADEMIC_LEVEL_CHOICES
    from courses.models import SEMESTER_CHOICES

    if request.method == 'POST' or 'all' in request.GET:
        dept  = request.POST.get('dept', '') if request.method == 'POST' else ''
        level = request.POST.get('level', '') if request.method == 'POST' else ''
        yr    = request.POST.get('year', '') if request.method == 'POST' else ''
        sem   = request.POST.get('sem', '') if request.method == 'POST' else ''

        qs = Result.objects.select_related('student', 'course').order_by('student__department', 'student__last_name', 'course__title')
        if dept:  qs = qs.filter(student__department=dept)
        if level: qs = qs.filter(student__academic_level=level)
        if yr:    qs = qs.filter(academic_year=yr)
        if sem:   qs = qs.filter(semester=sem)

        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from django.http import HttpResponse

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Results'
        header_font = Font(bold=True, color='FFFFFF', size=11)
        header_fill = PatternFill(start_color='1E4AAD', end_color='1E4AAD', fill_type='solid')
        thin_border = Border(
            left=Side(style='thin', color='D0D3E0'),
            right=Side(style='thin', color='D0D3E0'),
            top=Side(style='thin', color='D0D3E0'),
            bottom=Side(style='thin', color='D0D3E0'),
        )
        headers = ['#', 'Surname', 'First Name', 'Matric No', 'Department', 'Level',
                   'Course', 'Academic Year', 'Semester', 'CA /40', 'Exam /60', 'Total', 'Grade']
        ws.append(headers)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center')
            cell.border = thin_border

        grade_fills = {
            'A': PatternFill(start_color='D4EDDA', end_color='D4EDDA', fill_type='solid'),
            'B': PatternFill(start_color='D4EDDA', end_color='D4EDDA', fill_type='solid'),
            'C': PatternFill(start_color='FFF3CD', end_color='FFF3CD', fill_type='solid'),
            'D': PatternFill(start_color='FFF3CD', end_color='FFF3CD', fill_type='solid'),
            'E': PatternFill(start_color='FFF3CD', end_color='FFF3CD', fill_type='solid'),
            'F': PatternFill(start_color='F8D7DA', end_color='F8D7DA', fill_type='solid'),
        }

        for i, r in enumerate(qs, start=1):
            row_data = [
                i,
                r.student.last_name,
                r.student.first_name,
                r.student.matric_number or '',
                r.student.get_department_display() or '',
                r.student.get_academic_level_display() or '',
                r.course.title,
                r.academic_year,
                f'Semester {r.semester}',
                r.ca_score,
                r.exam_score,
                r.total,
                r.grade,
            ]
            ws.append(row_data)
            row_num = i + 1
            grade_fill = grade_fills.get(r.grade)
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=row_num, column=col_idx)
                cell.border = thin_border
                if grade_fill:
                    cell.fill = grade_fill

        for col in ws.columns:
            max_len = max((len(str(c.value or '')) for c in col), default=0)
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 44)

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        filename = f'results_{timezone.now():%Y%m%d_%H%M%S}.xlsx'
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        wb.save(response)
        return response

    dept_choices = [('', 'All Departments')] + list(get_department_choices())
    level_choices = [('', 'All Levels')] + list(ACADEMIC_LEVEL_CHOICES)
    semester_choices = [('', 'All Semesters')] + list(SEMESTER_CHOICES)
    year_choices = [('', 'All Years')] + [(y, y) for y in
                   Result.objects.values_list('academic_year', flat=True).distinct().order_by('-academic_year')]
    return render(request, 'portal/result_export.html', {
        'dept_choices': dept_choices,
        'level_choices': level_choices,
        'semester_choices': semester_choices,
        'year_choices': year_choices,
    })


@portal_required
def result_export_portal(request):
    """
    Push final Result rows from the LMS to the Portal (SMS) scores endpoint.
    Admin can export a whole year/semester or a single course, and optionally
    attribute the submission to a specific exam officer (staffId).
    """
    from api.sync import export_scores_to_portal
    from results.models import Result
    from courses.models import Course, SEMESTER_CHOICES
    from accounts.models import Student
    staff_choices = Student.objects.filter(
        role='lecturer', is_active=True, is_sms_synced=True
    ).order_by('last_name', 'first_name')

    if request.method == 'POST':
        yr    = request.POST.get('year', '').strip()
        sem   = request.POST.get('sem', '').strip()
        code  = request.POST.get('course', '').strip()
        staff = request.POST.get('staff_id', '').strip()
        if not yr or sem not in ('1', '2'):
            messages.error(request, 'Choose an academic year and a semester (1 or 2).')
        else:
            ok, info = export_scores_to_portal(
                yr, sem,
                staff_id=staff,
                course_code=code or None)
            if ok:
                messages.success(
                    request,
                    f'Exported {info["exported"]} result(s) for '
                    f'{yr} semester {sem}{f" · {code}" if code else ""} to the Portal.')
            else:
                messages.error(request, info['error'] or 'Scores export failed.')
        return redirect('portal:result_export_portal')

    year_choices = (Result.objects.values_list('academic_year', flat=True)
                    .distinct().order_by('-academic_year'))
    courses = (Course.objects.filter(status='published')
               .order_by('code'))
    return render(request, 'portal/result_export_portal.html', {
        'year_choices': year_choices,
        'semester_choices': SEMESTER_CHOICES,
        'course_choices': courses,
        'staff_choices': staff_choices,
    })


@portal_required
def result_edit(request, pk):
    from results.models import Result
    result = get_object_or_404(Result, pk=pk)
    if request.method == 'POST':
        result.ca_score     = float(request.POST.get('ca_score', result.ca_score))
        result.exam_score   = float(request.POST.get('exam_score', result.exam_score))
        result.ca_override  = True
        result.exam_override = True
        result.remark       = request.POST.get('remark', '')
        result.save()
        messages.success(request, 'Result updated.')
        return redirect('portal:result_list')
    return render(request, 'portal/result_edit.html', {'result': result})


@portal_required
def result_sync_all(request):
    from results.models import Result
    from courses.models import Enrollment
    count = 0
    for e in Enrollment.objects.select_related('student','course').filter(student__role='student'):
        Result.sync_for_student_course(e.student, e.course)
        count += 1
    messages.success(request, f'Synced {count} results.')
    return redirect('portal:result_list')


# ── Academic Settings ──────────────────────────────────────────────────

@portal_required
def academic_settings(request):
    from courses.models import AcademicSettings
    obj = AcademicSettings.get_current()
    if request.method == 'POST':
        obj.current_semester = request.POST.get('current_semester', obj.current_semester)
        obj.academic_year    = request.POST.get('academic_year', obj.academic_year)
        obj.save()
        messages.success(request, f'Updated: {obj.academic_year} Semester {obj.current_semester}')
        return redirect('portal:academic_settings')
    return render(request, 'portal/academic_settings.html', {'obj': obj})


# ── Announcements ──────────────────────────────────────────────────────

@portal_required
def announcement_list(request):
    from announcements.models import Announcement
    qs = Announcement.objects.select_related('created_by').all()
    q = request.GET.get('q', '').strip()
    if q:
        qs = qs.filter(Q(title__icontains=q) | Q(body__icontains=q))
    ctx = {
        'items': qs, 'q': q,
        'total_count': Announcement.objects.count(),
    }
    if _is_htmx(request):
        return render(request, 'portal/_announcement_table.html', ctx)
    return render(request, 'portal/announcement_list.html', ctx)


@portal_required
def announcement_create(request):
    from announcements.models import Announcement
    from accounts.models import get_department_choices, ACADEMIC_LEVEL_CHOICES
    if request.method == 'POST':
        p = request.POST
        depts  = ','.join(p.getlist('departments'))
        levels = ','.join(p.getlist('levels'))
        expires = p.get('expires_at') or None
        Announcement.objects.create(
            title=p['title'], body=p['body'],
            priority=p.get('priority','normal'),
            audience=p.get('audience','all'),
            departments=depts, levels=levels,
            expires_at=expires,
            created_by=request.user,
        )
        messages.success(request, 'Announcement posted.')
        return redirect('portal:announcement_list')
    return render(request, 'portal/announcement_form.html', {
        'dept_choices': get_department_choices(),
        'level_choices': ACADEMIC_LEVEL_CHOICES,
    })


@portal_required
def announcement_delete(request, pk):
    from announcements.models import Announcement
    ann = get_object_or_404(Announcement, pk=pk)
    if request.method == 'POST':
        ann.delete()
        messages.success(request, 'Announcement deleted.')
    return redirect('portal:announcement_list')


@portal_required
def announcement_toggle(request, pk):
    from announcements.models import Announcement
    ann = get_object_or_404(Announcement, pk=pk)
    if request.method == 'POST':
        ann.is_active = not ann.is_active
        ann.save(update_fields=['is_active'])
    return redirect('portal:announcement_list')


# ── Bulk Student Upload ────────────────────────────────────────────────

@portal_required
def student_bulk_upload(request):
    """
    Upload a CSV file to create multiple student accounts at once.
    Expected columns (with header row):
    first_name, last_name, email, matric_number, department, academic_level
    Password defaults to matric_number (or env DEFAULT_STUDENT_PASSWORD).
    """
    _DEFAULT_PWD = os.environ.get('DEFAULT_STUDENT_PASSWORD', 'changeme123')
    from accounts.models import Student, get_department_choices, ACADEMIC_LEVEL_CHOICES
    from courses.models import Enrollment
    import csv, io

    VALID_DEPTS  = {code for code,_ in get_department_choices()}
    VALID_LEVELS = {code for code,_ in ACADEMIC_LEVEL_CHOICES}

    results = None

    if request.method == 'POST' and request.FILES.get('csv_file'):
        csv_file = request.FILES['csv_file']
        try:
            text = csv_file.read().decode('utf-8-sig')  # handle BOM
        except UnicodeDecodeError:
            text = csv_file.read().decode('latin-1')

        reader = csv.DictReader(io.StringIO(text))
        # Normalise header names: strip whitespace, lowercase
        reader.fieldnames = [f.strip().lower().replace(' ','_') for f in (reader.fieldnames or [])]

        created, skipped, errors = [], [], []

        for i, row in enumerate(reader, start=2):  # row 1 = header
            row = {k: (v or '').strip() for k, v in row.items()}
            email        = row.get('email','')
            first_name   = row.get('first_name','')
            last_name    = row.get('last_name','')
            matric       = (row.get('matric_number','') or '').strip() or None
            department   = row.get('department','').upper()
            level        = row.get('academic_level','').upper()

            if not email or not first_name or not last_name:
                errors.append(f"Row {i}: missing first_name, last_name or email — skipped.")
                continue
            if Student.objects.filter(email=email).exists():
                skipped.append(f"Row {i}: {email} — account already exists.")
                continue
            if department and department not in VALID_DEPTS:
                errors.append(f"Row {i}: {email} — invalid department '{department}' — skipped.")
                continue
            if level and level not in VALID_LEVELS:
                errors.append(f"Row {i}: {email} — invalid level '{level}' — skipped.")
                continue

            password = matric or _DEFAULT_PWD
            s = Student.objects.create_user(
                username=email, email=email,
                first_name=first_name, last_name=last_name,
                password=password, role='student',
                department=department, academic_level=level,
                matric_number=matric,
            )
            Enrollment.sync_for_student(s)
            created.append(f"{s.get_full_name()} ({email})")

        results = {'created': created, 'skipped': skipped, 'errors': errors}
        if created:
            messages.success(request, f"{len(created)} student(s) created successfully.")

    return render(request, 'portal/student_bulk_upload.html', {'results': results})


# ── Exam Reschedule ────────────────────────────────────────────────────

@portal_required
def exam_reschedule_list(request, exam_pk):
    from exams.models import Exam, ExamReschedule
    from accounts.models import Student
    exam = get_object_or_404(Exam, pk=exam_pk)
    reschedules = ExamReschedule.objects.filter(exam=exam).select_related('student')
    # Students enrolled in this course who have NOT yet submitted an attempt
    from courses.models import Enrollment
    enrolled = Student.objects.filter(
        role='student',
        enrollments__course=exam.course,
    ).exclude(exam_attempts__exam=exam)
    return render(request, 'portal/exam_reschedule.html', {
        'exam': exam, 'reschedules': reschedules, 'enrolled_students': enrolled,
    })


@portal_required
def exam_reschedule_create(request, exam_pk):
    from exams.models import Exam, ExamReschedule
    exam = get_object_or_404(Exam, pk=exam_pk)
    if request.method == 'POST':
        from django.utils.dateparse import parse_datetime
        student_ids = request.POST.getlist('students')
        open_at  = parse_datetime(request.POST.get('open_at', ''))
        close_at = parse_datetime(request.POST.get('close_at', ''))
        reason   = request.POST.get('reason', '')
        if not open_at or not close_at:
            messages.error(request, 'Both open and close date/time are required.')
            return redirect('portal:exam_reschedule_list', exam_pk=exam_pk)
        if close_at <= open_at:
            messages.error(request, 'Close time must be after open time.')
            return redirect('portal:exam_reschedule_list', exam_pk=exam_pk)
        count = 0
        for sid in student_ids:
            ExamReschedule.objects.update_or_create(
                exam=exam, student_id=sid,
                defaults={'open_at': open_at, 'close_at': close_at, 'reason': reason}
            )
            count += 1
        messages.success(request, f'Rescheduled exam for {count} student(s).')
    return redirect('portal:exam_reschedule_list', exam_pk=exam_pk)


@portal_required
def exam_reschedule_delete(request, pk):
    from exams.models import ExamReschedule
    rs = get_object_or_404(ExamReschedule, pk=pk)
    exam_pk = rs.exam_id
    if request.method == 'POST':
        rs.delete()
        messages.success(request, 'Reschedule removed — student will use the global exam window.')
    return redirect('portal:exam_reschedule_list', exam_pk=exam_pk)


# ── Department Management ──────────────────────────────────────────────

@portal_required
def department_list(request):
    from departments.models import Department
    qs = Department.objects.all()
    q = request.GET.get('q', '').strip()
    if q:
        qs = qs.filter(Q(code__icontains=q) | Q(name__icontains=q))
    ctx = {
        'depts': qs, 'q': q,
        'total_count': Department.objects.count(),
    }
    if _is_htmx(request):
        return render(request, 'portal/_department_table.html', ctx)
    return render(request, 'portal/department_list.html', ctx)


@portal_required
def department_create(request):
    from departments.models import Department
    if request.method == 'POST':
        code = request.POST.get('code','').strip().upper()
        name = request.POST.get('name','').strip()
        order = request.POST.get('order', 0) or 0
        if not code or not name:
            messages.error(request, 'Code and name are required.')
        elif Department.objects.filter(code=code).exists():
            messages.error(request, f'Department with code "{code}" already exists.')
        else:
            Department.objects.create(code=code, name=name, order=order)
            messages.success(request, f'Department "{name}" ({code}) created.')
            return redirect('portal:department_list')
    return render(request, 'portal/department_form.html', {'is_new': True})


@portal_required
def department_edit(request, pk):
    from departments.models import Department
    dept = get_object_or_404(Department, pk=pk)
    if request.method == 'POST':
        # Code cannot be changed (it's the FK key used everywhere)
        dept.name      = request.POST.get('name', dept.name).strip()
        dept.order     = request.POST.get('order', dept.order) or 0
        dept.is_active = 'is_active' in request.POST
        dept.save()
        messages.success(request, f'Department "{dept.name}" updated.')
        return redirect('portal:department_list')
    return render(request, 'portal/department_form.html', {'is_new': False, 'dept': dept})


@portal_required
def course_message_list(request):
    """
    Admin oversight of every student→lecturer message thread — useful for courses
    with no lecturer assigned yet, or where a lecturer hasn't responded.
    """
    from courses.models import CourseMessage
    threads = CourseMessage.objects.select_related('course', 'student', 'course__lecturer')
    status_filter = request.GET.get('status')
    if status_filter in ('open', 'resolved'):
        threads = threads.filter(status=status_filter)
    return render(request, 'portal/course_message_list.html', {
        'threads': threads, 'status_filter': status_filter or 'all',
    })
